"""SPC (Annex A.2) Excel layout — mirrors the official Semi-Expendable Property Card PDF."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from equipment.coa import effective_cost
from equipment.spc import (
    get_spc_transactions,
    spc_fund_label,
    spc_lgu_name,
    spc_property_description,
    spc_semi_expendable_type,
    spc_useful_life_label,
)

from .coa_pdf import SPC_LEDGER_BODY_ROWS, SPC_LEDGER_HEADER, _spc_transaction_cells

_COLS = 13
_SIDE = Side(style='thin', color='000000')
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
_HEADER_FONT = Font(bold=True, size=8)
_LABEL_FONT = Font(size=8)
_TITLE_FONT = Font(bold=True, size=12)
_ANNEX_FONT = Font(size=9)

# Landscape column widths (approximate match to PDF proportions).
_COL_WIDTHS = [11, 12, 12, 9, 5, 10, 10, 9, 5, 18, 6, 10, 22]


def _set_col_widths(ws):
    for idx, width in enumerate(_COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _border_range(ws, row_start, row_end, col_start=1, col_end=_COLS):
    for row in range(row_start, row_end + 1):
        for col in range(col_start, col_end + 1):
            ws.cell(row, col).border = _BORDER


def _merge_write(ws, row, col_start, col_end, value, *, font, align, height=None):
    if col_start < col_end:
        ws.merge_cells(
            start_row=row, start_column=col_start,
            end_row=row, end_column=col_end,
        )
    cell = ws.cell(row, col_start, value)
    cell.font = font
    cell.alignment = align
    if height is not None:
        ws.row_dimensions[row].height = height


def _excel_sheet_title(property_number, index=0):
    """Sanitize property number for use as an Excel worksheet title (max 31 chars)."""
    invalid = set(r'\/*?:[]')
    cleaned = ''.join('-' if ch in invalid else ch for ch in (property_number or 'SPC')).strip()
    title = (cleaned or 'SPC')[:31]
    if index:
        suffix = f'-{index}'
        title = f'{title[:31 - len(suffix)]}{suffix}'
    return title or 'SPC'


def _write_ledger_header(ws, row):
    """Two-row SPC ledger header with merged group columns."""
    header1 = SPC_LEDGER_HEADER[0]
    header2 = SPC_LEDGER_HEADER[1]

    for col, label in enumerate(header1, start=1):
        cell = ws.cell(row, col, label.replace('\n', ' '))
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for col, label in enumerate(header2, start=1):
        cell = ws.cell(row + 1, col, label.replace('\n', ' '))
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    merges = [
        (row, row + 1, 1, 1),
        (row, row + 1, 2, 2),
        (row, row + 1, 3, 3),
        (row, row, 4, 7),
        (row, row, 8, 10),
        (row, row + 1, 11, 11),
        (row, row + 1, 12, 12),
        (row, row + 1, 13, 13),
    ]
    for r1, r2, c1, c2 in merges:
        if r1 != r2 or c1 != c2:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    _border_range(ws, row, row + 1)
    ws.row_dimensions[row].height = 24
    ws.row_dimensions[row + 1].height = 24


def _write_ledger_row(ws, row, values, *, data=False):
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.font = Font(size=8)
        cell.border = _BORDER
        if data:
            if col in (1, 2, 3, 4, 5, 8, 9, 11):
                cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
            elif col in (6, 7, 12):
                cell.alignment = Alignment(horizontal='right', vertical='top')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws.row_dimensions[row].height = 16


def _write_spc_form_sheet(ws, equipment):
    """Write one Annex A.2 Semi-Expendable Property Card worksheet."""
    _set_col_widths(ws)
    row = 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
    ws.merge_cells(start_row=row, start_column=12, end_row=row, end_column=13)
    ws.cell(row, 1, '')
    annex = ws.cell(row, 12, 'Annex A.2')
    annex.font = _ANNEX_FONT
    annex.alignment = Alignment(horizontal='right', vertical='top')
    ws.row_dimensions[row].height = 14
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=_COLS)
    title = ws.cell(row, 1, 'SEMI-EXPENDABLE PROPERTY CARD')
    title.font = _TITLE_FONT
    title.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 20
    row += 1

    _merge_write(
        ws, row, 1, 7, f'LGU: {spc_lgu_name()}',
        font=_LABEL_FONT, align=Alignment(vertical='bottom'), height=16,
    )
    _merge_write(
        ws, row, 8, 13, f'Fund: {spc_fund_label(equipment)}',
        font=_LABEL_FONT, align=Alignment(vertical='bottom'),
    )
    row += 1

    prop_type = spc_semi_expendable_type(equipment)
    useful_life = spc_useful_life_label(equipment)
    description = spc_property_description(equipment)
    prop_rows = [
        (f'Semi-expendable Property: {prop_type or ""}', f'Semi-expendable Property Number: {equipment.property_number}'),
        (f'Description: {description}', f'Estimated Useful Life: {useful_life}'),
    ]
    for left, right in prop_rows:
        _merge_write(
            ws, row, 1, 7, left,
            font=_LABEL_FONT, align=Alignment(vertical='top', wrap_text=True), height=18,
        )
        _merge_write(
            ws, row, 8, 13, right,
            font=_LABEL_FONT, align=Alignment(vertical='top', wrap_text=True),
        )
        _border_range(ws, row, row)
        row += 1

    row += 1
    table_start = row
    _write_ledger_header(ws, row)
    row += 2

    unit_cost = effective_cost(equipment)
    transactions = get_spc_transactions(equipment)
    data_rows = [_spc_transaction_cells(r, unit_cost) for r in transactions]
    for values in data_rows:
        _write_ledger_row(ws, row, values, data=True)
        row += 1

    blank_rows = max(0, SPC_LEDGER_BODY_ROWS - len(data_rows))
    for _ in range(blank_rows):
        _write_ledger_row(ws, row, [''] * _COLS)
        row += 1

    table_end = row - 1

    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4)
    ws.print_area = f'A{table_start - 4}:M{table_end}'


def generate_spc_excel(equipment):
    """Build an Excel workbook for one Semi-Expendable Property Card."""
    wb = Workbook()
    ws = wb.active
    ws.title = _excel_sheet_title(equipment.property_number)
    _write_spc_form_sheet(ws, equipment)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, f'SPC_{equipment.property_number}.xlsx'


def generate_all_spc_excel(equipment_list):
    """Build a workbook with one Annex A.2 sheet per equipment item."""
    items = list(equipment_list)
    if not items:
        raise ValueError('No semi-expendable equipment to export.')

    wb = Workbook()
    wb.remove(wb.active)
    used_titles = set()

    for index, equipment in enumerate(items, start=1):
        title = _excel_sheet_title(equipment.property_number)
        if title in used_titles:
            title = _excel_sheet_title(equipment.property_number, index)
        used_titles.add(title)
        ws = wb.create_sheet(title=title)
        _write_spc_form_sheet(ws, equipment)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, f'SPC_All_{len(items)}_items_{date.today()}.xlsx'
