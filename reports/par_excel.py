"""PAR (Appendix 71) Excel layout — mirrors the official PAR PDF form."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from core.models import SiteSettings
from core.signatories import active_supply_officer_user, default_supply_officer_user
from equipment.par import (
    get_active_par_for_equipment,
    par_entity_name,
    par_fund_cluster_display,
)

from .coa_pdf import (
    _fmt_amount,
    _group_pars_by_receipt,
    _manila_today,
    _par_compute_body_fill,
    _par_item_display_description,
    _par_line_items_from_equipment,
    _par_line_items_from_group,
    _par_max_items_per_page,
    _par_numbers_for_group,
    _par_sig_date_display,
    _par_signatories,
    _par_usable_dimensions,
)

_SIDE = Side(style='thin', color='000000')
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)

_COLS = 6
_ROW_HEADER = 22
_ROW_TABLE = 18
_ROW_SIG_HEADER = 20
_ROW_SIG_VALUE = 22
_ROW_SIG_NOTE = 16
_ROW_SIG_DATE_VAL = 16


def _set_height(ws, row, height):
    ws.row_dimensions[row].height = height


def _border_row(ws, row, cols=_COLS):
    for col in range(1, cols + 1):
        ws.cell(row, col).border = _BORDER


def _merge_row(ws, row, col_start, col_end, value, *, font, align, height):
    if col_start < col_end:
        ws.merge_cells(
            start_row=row, start_column=col_start,
            end_row=row, end_column=col_end,
        )
    cell = ws.cell(row, col_start, value)
    cell.font = font
    cell.alignment = align
    _border_row(ws, row)
    _set_height(ws, row, height)


def _write_item_row(ws, row, values, *, header=False):
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.border = _BORDER
        cell.font = Font(bold=header, size=9)
        if header:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        elif col in (1, 2, 5):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif col == 4:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if not header and len(str(value or '')) > 22:
                cell.font = Font(size=8)
        elif col == 6:
            cell.alignment = Alignment(horizontal='right', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    _set_height(ws, row, _ROW_TABLE)


def _write_sig_row(
    ws,
    row,
    left,
    right,
    *,
    left_font,
    right_font=None,
    left_align='left',
    right_align='left',
    height=_ROW_TABLE,
    pad_left=False,
    left_cols=3,
    right_cols=3,
):
    """Two-column signature row split across six columns."""
    if right_font is None:
        right_font = left_font
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=left_cols)
    ws.merge_cells(
        start_row=row, start_column=left_cols + 1, end_row=row, end_column=_COLS,
    )
    lc = ws.cell(row, 1, left)
    rc = ws.cell(row, left_cols + 1, right)
    lc.font = left_font
    rc.font = right_font
    lc.alignment = Alignment(
        horizontal=left_align, vertical='center', wrap_text=True,
        indent=1 if pad_left else 0,
    )
    rc.alignment = Alignment(
        horizontal=right_align, vertical='center', wrap_text=True,
        indent=1 if pad_left else 0,
    )
    _border_row(ws, row)
    _set_height(ws, row, height)


def _par_blank_row_count(item_count):
    _, usable_h = _par_usable_dimensions()
    blank_rows, _ = _par_compute_body_fill(item_count, usable_h)
    return blank_rows


def _setup_par_columns(ws, line_items=None):
    ws.column_dimensions['A'].width = 9
    ws.column_dimensions['B'].width = 8
    desc_width = 34
    prop_width = 16
    if line_items:
        max_prop = max(
            (len(str(getattr(item, 'property_number', '') or '')) for item in line_items),
            default=0,
        )
        if max_prop > 16:
            prop_width = min(30, max(16, max_prop * 0.82 + 2))
            desc_width = max(26, 34 - (prop_width - 16) * 0.85)
    ws.column_dimensions['C'].width = desc_width
    ws.column_dimensions['D'].width = prop_width
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12


def _excel_sheet_title(par_no, index=0):
    """Sanitize PAR number for use as an Excel worksheet title (max 31 chars)."""
    invalid = set(r'\/*?:[]')
    cleaned = ''.join('-' if ch in invalid else ch for ch in (par_no or 'PAR')).strip()
    title = (cleaned or 'PAR')[:31]
    if index:
        suffix = f'-{index}'
        title = f'{title[:31 - len(suffix)]}{suffix}'
    return title or 'PAR'


def _write_par_form_sheet(
    ws,
    *,
    site,
    par_no,
    entity_name,
    fund_cluster,
    line_items,
    signatories,
    sig_date_by,
    sig_date_from=None,
    page_label='Page 1/1',
):
    """Write one Appendix 71 PAR form onto a worksheet."""
    if sig_date_from is None:
        sig_date_from = sig_date_by

    _setup_par_columns(ws, line_items)

    row = 1
    annex = ws.cell(row, 6, 'Appendix 71')
    annex.font = Font(italic=True, size=9)
    annex.alignment = Alignment(horizontal='right', vertical='bottom')
    _set_height(ws, row, 15)
    row += 1

    _merge_row(
        ws, row, 1, 6, site.par_header_title.upper(),
        font=Font(bold=True, size=12),
        align=Alignment(horizontal='center', vertical='center'),
        height=_ROW_HEADER,
    )
    row += 1

    _merge_row(
        ws, row, 1, 6, f'Entity Name: {entity_name}',
        font=Font(size=10),
        align=Alignment(horizontal='left', vertical='center', indent=1),
        height=18,
    )
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
    fund_cell = ws.cell(row, 1, f'Fund Cluster: {fund_cluster}')
    fund_cell.font = Font(size=10)
    fund_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    par_cell = ws.cell(row, 5, f'PAR No.: {par_no}')
    par_cell.font = Font(size=10)
    par_cell.alignment = Alignment(horizontal='right', vertical='center')
    _border_row(ws, row)
    _set_height(ws, row, 18)
    row += 1

    row += 1
    table_start = row
    _write_item_row(
        ws, row,
        ['Quantity', 'Unit', 'Description', 'Property Number', 'Date Acquired', 'Amount'],
        header=True,
    )
    row += 1

    for item in line_items:
        _write_item_row(ws, row, [
            item.quantity,
            item.unit,
            _par_item_display_description(item),
            item.property_number,
            str(item.date_acquired or ''),
            _fmt_amount(item.amount),
        ])
        row += 1

    for _ in range(_par_blank_row_count(len(line_items))):
        _write_item_row(ws, row, ['', '', '', '', '', ''])
        row += 1

    sig_header = Font(bold=True, size=9)
    sig_note = Font(size=7)
    sig_name = Font(bold=True, size=9, underline='single')
    sig_pos = Font(bold=True, size=9, underline='single')
    sig_date_font = Font(bold=True, size=8, underline='single')

    _write_sig_row(
        ws, row, 'Received by:', 'Issued by:',
        left_font=sig_header, left_align='left', right_align='left',
        height=_ROW_SIG_HEADER, pad_left=True,
    )
    row += 1

    _write_sig_row(
        ws, row,
        (signatories['by_name'] or '').upper(),
        (signatories['from_name'] or '').upper(),
        left_font=sig_name, left_align='center', right_align='center',
        height=_ROW_SIG_VALUE,
    )
    row += 1

    _write_sig_row(
        ws, row,
        'Signature over Printed Name of End User',
        'Signature over Printed Name of Supply and/or Property Custodian',
        left_font=sig_note, left_align='center', right_align='center',
        height=_ROW_SIG_NOTE,
    )
    row += 1

    _write_sig_row(
        ws, row,
        (signatories['by_position'] or '').upper(),
        (signatories['from_position'] or '').upper(),
        left_font=sig_pos, left_align='center', right_align='center',
        height=_ROW_SIG_VALUE,
    )
    row += 1

    _write_sig_row(
        ws, row,
        'Position/Office',
        'Position/Office',
        left_font=sig_note, left_align='center', right_align='center',
        height=_ROW_SIG_NOTE,
    )
    row += 1

    _write_sig_row(
        ws, row,
        _par_sig_date_display(sig_date_by),
        _par_sig_date_display(sig_date_from),
        left_font=sig_date_font, left_align='center', right_align='center',
        height=_ROW_SIG_DATE_VAL,
    )
    row += 1

    _write_sig_row(
        ws, row, 'Date', 'Date',
        left_font=sig_note, left_align='center', right_align='center',
        height=_ROW_SIG_NOTE,
    )
    row += 1

    form_end = row - 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
    footer_left = ws.cell(row, 1, site.par_header_title.title())
    footer_left.font = Font(italic=True, size=8, color='808080')
    footer_left.alignment = Alignment(horizontal='left', vertical='center')
    footer_right = ws.cell(row, 4, page_label)
    footer_right.font = Font(italic=True, size=8, color='808080')
    footer_right.alignment = Alignment(horizontal='right', vertical='center')
    _set_height(ws, row, 14)

    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.4, bottom=0.4)
    ws.print_area = f'A{table_start}:F{form_end}'


def generate_par_excel(
    equipment_list, supply_officer=None, end_user=None, signatory_overrides=None,
    *, entity_name=None, fund_cluster=None, filename=None,
):
    """Build an Excel workbook matching the PAR PDF (Appendix 71) layout."""
    equipment = list(equipment_list)
    if not equipment:
        raise ValueError('No equipment to export.')

    site = SiteSettings.get_solo()
    officer = (
        active_supply_officer_user(user=supply_officer)
        or default_supply_officer_user()
    )
    if end_user is None:
        from core.signatories import resolve_par_received_by_user
        end_user = resolve_par_received_by_user(equipment_list=equipment)
    if not entity_name:
        entity_name = par_entity_name(equipment_list=equipment, end_user=end_user)
    if not fund_cluster:
        fund_cluster = par_fund_cluster_display(equipment_list=equipment)

    line_items = _par_line_items_from_equipment(equipment)
    signatories = _par_signatories(
        site,
        supply_officer=officer,
        end_user=end_user,
        signatory_overrides=signatory_overrides,
    )
    sig_date = str(_manila_today())
    par_no = ''
    for eq in equipment:
        active_par = get_active_par_for_equipment(eq)
        if active_par:
            par_no = active_par.par_no
            break

    wb = Workbook()
    ws = wb.active
    ws.title = _excel_sheet_title(par_no)
    _write_par_form_sheet(
        ws,
        site=site,
        par_no=par_no,
        entity_name=entity_name,
        fund_cluster=fund_cluster,
        line_items=line_items,
        signatories=signatories,
        sig_date_by=sig_date,
        sig_date_from=sig_date,
    )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, f'PAR_Report_{len(equipment)}_items_{date.today()}.xlsx'


def generate_all_par_excel(par_list, supply_officer=None, end_user=None, signatory_overrides=None):
    """Build a workbook — same entity, fund cluster, and received-by on one sheet."""
    par_records = list(par_list)
    if not par_records:
        raise ValueError('No PAR records to export.')

    site = SiteSettings.get_solo()
    officer = (
        active_supply_officer_user(user=supply_officer)
        or default_supply_officer_user()
    )
    wb = Workbook()
    wb.remove(wb.active)

    used_titles = set()
    sheet_index = 0
    per_page = _par_max_items_per_page()
    for group in _group_pars_by_receipt(par_records):
        primary = group[0]
        all_line_items = _par_line_items_from_group(group)
        par_no = _par_numbers_for_group(group) if len(group) > 1 else primary.par_no
        entity_name = par_entity_name(par=primary)
        fund_cluster = par_fund_cluster_display(par=primary)
        group_end_user = end_user or primary.end_user
        signatories = _par_signatories(
            site,
            par=primary,
            supply_officer=officer,
            end_user=group_end_user,
            signatory_overrides=signatory_overrides,
        )
        chunks = [
            all_line_items[i:i + per_page]
            for i in range(0, len(all_line_items), per_page)
        ] or [[]]
        page_total = len(chunks)
        for idx, chunk in enumerate(chunks, start=1):
            sheet_index += 1
            title = _excel_sheet_title(par_no)
            if title in used_titles:
                title = _excel_sheet_title(par_no, sheet_index)
            used_titles.add(title)

            ws = wb.create_sheet(title=title)
            _write_par_form_sheet(
                ws,
                site=site,
                par_no=par_no,
                entity_name=entity_name,
                fund_cluster=fund_cluster,
                line_items=chunk,
                signatories=signatories,
                sig_date_by=str(primary.date_received),
                sig_date_from=str(primary.date_issued),
                page_label=f'Page {idx}/{page_total}',
            )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    item_count = sum(par.items.count() for par in par_records)
    return output, f'PAR_All_{item_count}_items_{date.today()}.xlsx'
