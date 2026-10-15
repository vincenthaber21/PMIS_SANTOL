"""Appendix 70 Property, Plant and Equipment Ledger Card — Excel layout.

Matches the municipal Depreciation workbook: landscape ledger, receipt columns,
annual depreciation rows, and repair history beside the adjusted cost.
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from equipment.ppe_ledger import DASH, build_ppe_ledger

_FONT = 'Times New Roman'
_THIN = Side(style='thin', color='000000')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY = '#,##0.00'
_DATE = 'mm-dd-yy'

_COL_WIDTHS = {
    'A': 3,
    'B': 14,
    'C': 16,
    'D': 10,
    'E': 14,
    'F': 15,
    'G': 15,
    'H': 16,
    'I': 16,
    'J': 16,
    'K': 16,
    'L': 20,
    'M': 14,
}


def _font(size, *, bold=False, italic=False):
    return Font(name=_FONT, size=size, bold=bold, italic=italic)


def _align(horizontal='center', *, wrap=False, vertical='center'):
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


def _excel_sheet_title(property_number, index=0):
    invalid = set(r'\/*?:[]')
    cleaned = ''.join('-' if ch in invalid else ch for ch in (property_number or 'PPELC')).strip()
    title = (cleaned or 'PPELC')[:31]
    if index:
        suffix = f'-{index}'
        title = f'{title[:31 - len(suffix)]}{suffix}'
    return title or 'PPELC'


def _safe_filename(property_number):
    invalid = set('\\/:*?"<>|')
    cleaned = ''.join('-' if ch in invalid else ch for ch in (property_number or 'PPELC'))
    return cleaned.strip() or 'PPELC'


def _set_columns(ws):
    for letter, width in _COL_WIDTHS.items():
        ws.column_dimensions[letter].width = width


def _merge(ws, row, col, end_row, end_col, value, *, font, align, height=None):
    if row != end_row or col != end_col:
        ws.merge_cells(start_row=row, start_column=col, end_row=end_row, end_column=end_col)
    cell = ws.cell(row, col, value)
    cell.font = font
    cell.alignment = align
    if height is not None:
        for line in range(row, end_row + 1):
            ws.row_dimensions[line].height = height
    return cell


def _paint(cell, value, *, font, align, border=False, money=False, as_date=False):
    cell.font = font
    cell.alignment = align
    if border:
        cell.border = _BORDER
    if value is None or value == '':
        cell.value = None
        return
    if value == DASH:
        cell.value = DASH
        return
    if as_date:
        cell.value = value
        cell.number_format = _DATE
        return
    if money:
        cell.value = float(value)
        cell.number_format = _MONEY
        return
    cell.value = value


def _border_row(ws, row, start=2, end=13):
    for col in range(start, end + 1):
        ws.cell(row, col).border = _BORDER


def _write_table_header(ws, row):
    header_font = _font(11, bold=True)
    center = _align('center', wrap=True)
    labels = {
        (row, 2): 'Date',
        (row, 3): 'Reference',
        (row, 4): 'Receipt',
        (row, 8): 'Accumulated',
        (row, 9): 'Accumulated',
        (row, 10): 'Issues/Transfers/',
        (row, 11): 'Adjusted',
        (row, 12): 'Repair History',
        (row + 1, 4): 'Qty.',
        (row + 1, 5): 'Unit Cost',
        (row + 1, 6): 'Total Cost',
        (row + 1, 7): 'Depreciation',
        (row + 1, 8): 'Depreciation',
        (row + 1, 9): 'Impairment Losses',
        (row + 1, 10): 'Adjustment/s',
        (row + 1, 11): 'Cost',
        (row + 1, 12): 'Nature of Repair',
        (row + 1, 13): 'Amount',
    }
    merges = [
        (row, 2, row + 1, 2),
        (row, 3, row + 1, 3),
        (row, 4, row, 6),
        (row, 12, row, 13),
    ]
    for start_row, start_col, end_row, end_col in merges:
        ws.merge_cells(
            start_row=start_row, start_column=start_col,
            end_row=end_row, end_column=end_col,
        )
    for (line, col), label in labels.items():
        cell = ws.cell(line, col, label)
        cell.font = header_font
        cell.alignment = center
    for line in (row, row + 1):
        _border_row(ws, line)
        ws.row_dimensions[line].height = 32
        for col in range(2, 14):
            ws.cell(line, col).font = header_font
            ws.cell(line, col).alignment = center
            ws.cell(line, col).border = _BORDER


def _write_data_row(ws, row, entry):
    body = _font(11)
    center = _align('center', wrap=True)
    values = [
        (2, entry.get('date'), True, False),
        (3, entry.get('reference'), False, False),
        (4, entry.get('qty'), False, False),
        (5, entry.get('unit_cost'), False, True),
        (6, entry.get('total_cost'), False, True),
        (7, entry.get('depreciation'), False, True),
        (8, entry.get('accumulated'), False, True),
        (9, entry.get('impairment'), False, True),
        (10, entry.get('adjustment'), False, True),
        (11, entry.get('adjusted_cost'), False, True),
        (12, entry.get('repair_nature'), False, False),
        (13, entry.get('repair_amount'), False, True),
    ]
    for col, value, as_date, money in values:
        if value == DASH:
            money = False
        _paint(
            ws.cell(row, col),
            value,
            font=body,
            align=center,
            border=True,
            money=money and value not in (None, DASH),
            as_date=as_date and value not in (None, DASH),
        )
    ws.row_dimensions[row].height = 18


def _apply_page(ws, last_row):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.45, header=0.2, footer=0.2)
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:M{last_row}'


def write_ppe_ledger_sheet(ws, ledger):
    """Write one Appendix 70 card onto ``ws``."""
    _set_columns(ws)
    title = _font(14, bold=True)
    heading = _font(12, bold=True)
    body = _font(12)
    italic = _font(14, italic=True)
    center = _align('center', wrap=True)
    left = _align('left', wrap=True, vertical='center')
    right = _align('right', wrap=True, vertical='center')

    _merge(ws, 3, 12, 3, 13, ledger.get('appendix') or 'Appendix 70', font=italic, align=right)

    _merge(ws, 8, 2, 8, 13, ledger.get('republic') or 'Republic of the Philippines', font=body, align=center, height=18)
    _merge(ws, 9, 2, 9, 13, ledger.get('province') or '', font=body, align=center, height=18)
    _merge(ws, 10, 2, 10, 13, ledger.get('municipality') or '', font=heading, align=center, height=18)
    _merge(
        ws, 13, 2, 13, 13,
        ledger.get('title') or 'PROPERTY, PLANT AND EQUIPMENT LEDGER CARD',
        font=title, align=center, height=22,
    )

    office = ledger.get('accountable_office') or ''
    fund = ledger.get('fund_cluster') or ''
    officer = ledger.get('accountable_officer') or ''
    accountability = ledger.get('accountability_date') or ''
    _merge(ws, 16, 2, 16, 9, f'Accountable Office : {office}', font=heading, align=left, height=18)
    _merge(ws, 16, 11, 16, 13, f'Fund Cluster : {fund}', font=heading, align=left)
    _merge(ws, 17, 2, 17, 10, f'Accountable Officer: {officer}', font=heading, align=left, height=18)
    _merge(ws, 17, 12, 17, 13, f'Date of Accountability: {accountability}', font=heading, align=right)

    _merge(
        ws, 18, 2, 19, 10,
        f"Property, Plant and Equipment: {ledger.get('ppe_type') or ''}",
        font=body, align=left, height=18,
    )
    _merge(ws, 18, 11, 18, 13, f"Object Account Code: {ledger.get('object_code') or ''}", font=heading, align=left)
    _merge(ws, 19, 11, 19, 13, f"Estimated Useful Life:  {ledger.get('useful_life_label') or ''}", font=heading, align=left)

    description_height = 22 if len(ledger.get('description') or '') > 90 else 18
    _merge(
        ws, 20, 2, 21, 10,
        f"Description:  {ledger.get('description') or ''}",
        font=body, align=left, height=description_height,
    )
    _merge(ws, 20, 11, 20, 13, f"Rate of Depreciation:   {ledger.get('rate_label') or ''}", font=heading, align=left)
    _merge(ws, 21, 11, 21, 13, f"Remarks: {ledger.get('remarks') or ''}", font=heading, align=left)

    header_row = 22
    _write_table_header(ws, header_row)
    row = header_row + 2
    for entry in ledger.get('rows') or []:
        _write_data_row(ws, row, entry)
        row += 1
    if not ledger.get('rows'):
        _border_row(ws, row)
        ws.row_dimensions[row].height = 18
        row += 1

    row += 2
    label_font = _font(12, bold=True)
    name_font = _font(12, bold=True)
    pos_font = _font(12)
    _merge(ws, row, 2, row, 6, 'PREPARED BY:', font=label_font, align=_align('left', vertical='center'))
    _merge(ws, row, 11, row, 13, 'NOTED BY:', font=label_font, align=_align('left', vertical='center'))
    row += 3
    _merge(ws, row, 2, row, 6, (ledger.get('prepared_by_name') or '').upper(), font=name_font, align=center, height=18)
    _merge(ws, row, 11, row, 13, (ledger.get('noted_by_name') or '').upper(), font=name_font, align=center)
    row += 1
    _merge(ws, row, 2, row, 6, ledger.get('prepared_by_position') or '', font=pos_font, align=center, height=18)
    _merge(ws, row, 11, row, 13, ledger.get('noted_by_position') or '', font=pos_font, align=center)

    percent = ledger.get('depreciable_percent')
    percent_text = format(percent, 'f').rstrip('0').rstrip('.') if percent is not None else '95'
    row += 2
    _merge(
        ws, row, 2, row, 13,
        (
            f'Depreciable value: {percent_text}% of acquisition cost. '
            'A purchase on days 1-15 starts depreciation in that month; '
            'day 16 or later starts on the first day of the next month.'
        ),
        font=_font(9, italic=True),
        align=_align('left', wrap=True, vertical='center'),
        height=28,
    )
    _apply_page(ws, row)
    return row


def generate_ppe_ledger_excel(equipment, *, site=None):
    """Build a workbook with one Appendix 70 sheet for a single PPE item."""
    if site is None:
        from core.models import SiteSettings
        site = SiteSettings.get_solo()
    ledger = build_ppe_ledger(equipment, site=site)
    wb = Workbook()
    ws = wb.active
    ws.title = _excel_sheet_title(ledger.get('property_number'))
    write_ppe_ledger_sheet(ws, ledger)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"PPELC_{_safe_filename(ledger.get('property_number'))}.xlsx"
    return output, filename


def generate_all_ppe_ledger_excel(equipment_list, *, site=None):
    """Build a workbook with one Appendix 70 sheet per PPE item."""
    items = list(equipment_list)
    if not items:
        raise ValueError('No PPE equipment to export.')
    if site is None:
        from core.models import SiteSettings
        site = SiteSettings.get_solo()

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for index, equipment in enumerate(items, start=1):
        ledger = build_ppe_ledger(equipment, site=site)
        title = _excel_sheet_title(ledger.get('property_number'))
        if title in used:
            title = _excel_sheet_title(ledger.get('property_number'), index)
        used.add(title)
        ws = wb.create_sheet(title=title)
        write_ppe_ledger_sheet(ws, ledger)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, f'PPE_Ledger_Cards_{len(items)}_items_{date.today().isoformat()}.xlsx'
