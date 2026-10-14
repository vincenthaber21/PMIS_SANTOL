"""Reports views — generate and export PDF, Excel, CSV."""
import csv
import io
from collections import defaultdict
from datetime import date
from django.shortcuts import render
from core.decorators import insights_login_required
from django.contrib.auth import get_user_model
from django.http import HttpResponse, Http404
from django.db.models import Count, F, Q
from equipment.models import Equipment, Department, Category, FundSource

User = get_user_model()

REPORT_TYPES = [
    ('property_all', 'All Property (by Office/Personnel, Category & Fund Source)'),
    ('property_working', 'Serviceable / Working Property'),
    ('property_unserviceable', 'Unserviceable Property'),
    ('property_disposed', 'Disposed Property'),
    ('inventory', 'Inventory Report'),
    ('asset_valuation', 'Asset Valuation Report'),
    ('performance', 'Equipment Performance Report'),
]

PROPERTY_REPORT_TYPES = {
    'property_all', 'property_working', 'property_unserviceable', 'property_disposed',
}

PROPERTY_COLUMNS = [
    'Property No.', 'Name', 'Category', 'Fund Source',
    'Office', 'Personnel', 'Property Status', 'Brand', 'Location', 'Acquisition Cost (₱)',
]

# property_working uses exclusion logic (see _property_queryset) so items with no
# property_status yet still count as serviceable, matching new-equipment defaults.
PROPERTY_STATUS_FILTER = {
    'property_all': None,
    'property_working': 'serviceable',
    'property_unserviceable': 'Unserviceable',
    'property_disposed': 'Disposed',
}


def _personnel_name(user):
    if not user:
        return 'Unassigned'
    return user.get_full_name() or user.username


DATE_ORDER_CHOICES = [
    ('', 'Default (Office / Personnel)'),
    ('newest', 'Newest to Oldest'),
    ('oldest', 'Oldest to Newest'),
]


def _parse_filters(params):
    """Extract common report filter values from GET parameters."""
    location_id = params.get('location', '') or params.get('department', '') or None
    date_order = params.get('date_order', '') or ''
    if date_order not in ('newest', 'oldest'):
        date_order = ''
    return {
        'dept_id': location_id,
        'location_id': location_id,
        'category_id': params.get('category', '') or None,
        'fund_source_id': params.get('fund_source', '') or None,
        'personnel_id': params.get('personnel', '') or None,
        'status': params.get('status', '') or None,
        'date_order': date_order,
    }


def _date_order_by(date_order, *, grouped=False):
    """Order_by fields for purchase date (fallback: record created date)."""
    if date_order == 'newest':
        date_fields = [
            F('purchase_date').desc(nulls_last=True),
            F('created_at').desc(nulls_last=True),
        ]
    elif date_order == 'oldest':
        date_fields = [
            F('purchase_date').asc(nulls_last=True),
            F('created_at').asc(nulls_last=True),
        ]
    else:
        return []

    if grouped:
        return [
            'department__name',
            'assigned_user__last_name',
            'assigned_user__first_name',
            'category__name',
            'fund_source__name',
            *date_fields,
            'property_number',
        ]
    return [*date_fields, 'property_number']


def _property_queryset(report_type, filters):
    """Base queryset for property reports with shared filters."""
    qs = Equipment.objects.select_related(
        'department', 'assigned_user', 'category', 'fund_source', 'property_status',
    )
    date_order = filters.get('date_order', '')
    order_by = _date_order_by(date_order, grouped=True) or [
        'department__name',
        'assigned_user__last_name',
        'assigned_user__first_name',
        'category__name',
        'fund_source__name',
        'property_number',
    ]
    qs = qs.order_by(*order_by)

    status_name = PROPERTY_STATUS_FILTER.get(report_type)
    if status_name == 'serviceable':
        qs = qs.exclude(
            property_status__name__iexact='Unserviceable',
        ).exclude(
            property_status__name__iexact='Disposed',
        ).exclude(
            property_status__name__iexact='Condemned',  # legacy
        )
    elif status_name == 'Disposed':
        qs = qs.filter(
            Q(property_status__name__iexact='Disposed')
            | Q(property_status__name__iexact='Condemned')  # legacy
        )
    elif status_name:
        qs = qs.filter(property_status__name__iexact=status_name)

    if filters.get('dept_id'):
        qs = qs.filter(department_id=filters['dept_id'])
    if filters.get('category_id'):
        qs = qs.filter(category_id=filters['category_id'])
    if filters.get('fund_source_id'):
        qs = qs.filter(fund_source_id=filters['fund_source_id'])
    if filters.get('personnel_id'):
        qs = qs.filter(assigned_user_id=filters['personnel_id'])

    return qs


def property_row_for(obj):
    """Serialize an Equipment instance for property reports."""
    return [
        obj.property_number,
        obj.display_name,
        obj.category.name if obj.category_id else '—',
        obj.fund_source.name if obj.fund_source_id else (obj.fund or '—'),
        obj.department.name if obj.department_id else '—',
        _personnel_name(obj.assigned_user),
        obj.property_status.name if obj.property_status_id else '—',
        obj.brand,
        obj.display_location,
        float(obj.acquisition_cost or obj.purchase_cost or 0),
    ]


def group_property_data(qs):
    """Group equipment rows by Office → Personnel → Category → Fund Source."""
    tree = defaultdict(
        lambda: defaultdict(
            lambda: defaultdict(
                lambda: defaultdict(list)
            )
        )
    )
    for obj in qs:
        office = obj.department.name if obj.department_id else 'Unassigned Office'
        personnel = _personnel_name(obj.assigned_user)
        category = obj.category.name if obj.category_id else 'Uncategorized'
        fund = obj.fund_source.name if obj.fund_source_id else (obj.fund or 'Unspecified Fund')
        tree[office][personnel][category][fund].append(property_row_for(obj))

    sections = []
    for office in sorted(tree.keys()):
        for personnel in sorted(tree[office].keys()):
            for category in sorted(tree[office][personnel].keys()):
                for fund in sorted(tree[office][personnel][category].keys()):
                    rows = tree[office][personnel][category][fund]
                    sections.append({
                        'office': office,
                        'personnel': personnel,
                        'category': category,
                        'fund_source': fund,
                        'rows': rows,
                        'count': len(rows),
                    })
    return sections


def get_report_data(report_type, filters=None):
    """Return queryset and columns for the given report type."""
    filters = filters or {}

    if report_type in PROPERTY_REPORT_TYPES:
        return _property_queryset(report_type, filters), PROPERTY_COLUMNS

    dept_id = filters.get('dept_id')
    category_id = filters.get('category_id')
    fund_source_id = filters.get('fund_source_id')
    personnel_id = filters.get('personnel_id')
    status = filters.get('status')
    date_order = filters.get('date_order', '')
    date_order_by = _date_order_by(date_order) or ['property_number']

    if report_type == 'inventory':
        qs = Equipment.objects.select_related('department', 'assigned_user', 'category')
        if dept_id:
            qs = qs.filter(department_id=dept_id)
        if category_id:
            qs = qs.filter(category_id=category_id)
        if fund_source_id:
            qs = qs.filter(fund_source_id=fund_source_id)
        if personnel_id:
            qs = qs.filter(assigned_user_id=personnel_id)
        if status:
            qs = qs.filter(status=status)
        return qs.order_by(*date_order_by), ['Property No.', 'Name', 'Category', 'Brand',
                    'Department', 'Status', 'Condition', 'Location']

    if report_type == 'asset_valuation':
        qs = Equipment.objects.select_related('department').filter(status__in=['Available', 'Borrowed'])
        if dept_id:
            qs = qs.filter(department_id=dept_id)
        if category_id:
            qs = qs.filter(category_id=category_id)
        if fund_source_id:
            qs = qs.filter(fund_source_id=fund_source_id)
        if personnel_id:
            qs = qs.filter(assigned_user_id=personnel_id)
        return qs.order_by(*date_order_by), ['Property No.', 'Name', 'Purchase Cost (₱)', 'Book Value (₱)',
                    'Age (yrs)', 'Depreciation Method', 'Department']

    if report_type == 'performance':
        qs = Equipment.objects.select_related('department').annotate(
            borrow_count=Count('borrow_requests'),
            repair_count=Count('maintenance_requests'),
        )
        if dept_id:
            qs = qs.filter(department_id=dept_id)
        if category_id:
            qs = qs.filter(category_id=category_id)
        if fund_source_id:
            qs = qs.filter(fund_source_id=fund_source_id)
        if personnel_id:
            qs = qs.filter(assigned_user_id=personnel_id)
        return qs.order_by(*date_order_by), ['Property No.', 'Name', 'Utilization %', 'Condition',
                    'Borrow Count', 'Repair Count', 'Risk Score']

    return Equipment.objects.none(), []


def row_for(obj, report_type):
    """Convert a model instance to a CSV/Excel row for the given report type."""
    if report_type in PROPERTY_REPORT_TYPES:
        return property_row_for(obj)
    if report_type == 'inventory':
        return [obj.property_number, obj.display_name, obj.category.name,
                obj.brand, obj.department.name if obj.department else '', obj.status,
                obj.condition_score, obj.display_location]
    if report_type == 'asset_valuation':
        return [obj.property_number, obj.display_name, float(obj.purchase_cost),
                float(obj.current_book_value), obj.age_years,
                obj.get_depreciation_method_display(),
                obj.department.name if obj.department else '']
    if report_type == 'performance':
        return [obj.property_number, obj.display_name, float(obj.utilization_percent),
                obj.condition_score, obj.borrow_count, obj.repair_count, obj.risk_score]
    return []


def _excel_sheet_title(label):
    """Sanitize a string for use as an Excel worksheet title (max 31 chars)."""
    invalid = set(r'\/*?:[]')
    cleaned = ''.join('-' if ch in invalid else ch for ch in label).strip()
    return (cleaned or 'Report')[:31]


def _export_query_string(request):
    """Preserve current filter parameters for export links."""
    return request.GET.urlencode()


def _report_context(request):
    """Build shared context for report index and exports."""
    selected_type = request.GET.get('type', 'property_all')
    filters = _parse_filters(request.GET)
    qs, columns = get_report_data(selected_type, filters)

    sections = []
    rows = []
    if selected_type in PROPERTY_REPORT_TYPES:
        sections = group_property_data(qs)
        for section in sections:
            rows.extend(section['rows'])
    else:
        rows = [row_for(obj, selected_type) for obj in qs]

    return {
        'selected_type': selected_type,
        'filters': filters,
        'qs': qs,
        'columns': columns,
        'rows': rows,
        'sections': sections,
        'count': len(rows),
        'export_qs': _export_query_string(request),
    }


@insights_login_required
def reports_index(request):
    """Reports generation page."""
    ctx = _report_context(request)
    personnel = User.objects.filter(
        assigned_equipment__isnull=False,
    ).distinct().order_by('last_name', 'first_name')

    return render(request, 'reports/index.html', {
        'report_types': REPORT_TYPES,
        'selected_type': ctx['selected_type'],
        'departments': Department.objects.filter(status='Active'),
        'categories': Category.objects.filter(status='Active'),
        'fund_sources': FundSource.objects.filter(status='Active'),
        'personnel_list': personnel,
        'location_id': ctx['filters'].get('location_id') or '',
        'dept_id': ctx['filters'].get('dept_id') or '',
        'category_id': ctx['filters'].get('category_id') or '',
        'fund_source_id': ctx['filters'].get('fund_source_id') or '',
        'personnel_id': ctx['filters'].get('personnel_id') or '',
        'status': ctx['filters'].get('status') or '',
        'date_order': ctx['filters'].get('date_order') or '',
        'date_order_choices': DATE_ORDER_CHOICES,
        'columns': ctx['columns'],
        'rows': ctx['rows'],
        'sections': ctx['sections'],
        'count': ctx['count'],
        'export_qs': ctx['export_qs'],
        'is_property_report': ctx['selected_type'] in PROPERTY_REPORT_TYPES,
        'report_generated': bool(request.GET),
    })


@insights_login_required
def export_csv(request):
    """Export report as CSV."""
    ctx = _report_context(request)
    report_type = ctx['selected_type']
    qs, columns = get_report_data(report_type, ctx['filters'])

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{report_type}_{date.today()}.csv"'
    writer = csv.writer(response)
    writer.writerow(columns)

    if report_type in PROPERTY_REPORT_TYPES:
        for section in ctx['sections']:
            writer.writerow([])
            writer.writerow([
                f"Office: {section['office']}",
                f"Personnel: {section['personnel']}",
                f"Category: {section['category']}",
                f"Fund Source: {section['fund_source']}",
            ])
            for row in section['rows']:
                writer.writerow(row)
    else:
        for obj in qs:
            writer.writerow(row_for(obj, report_type))
    return response


@insights_login_required
def export_excel(request):
    """Export report as Excel file."""
    ctx = _report_context(request)
    report_type = ctx['selected_type']

    if report_type in PROPERTY_REPORT_TYPES:
        filters = ctx['filters']
        qs = _property_queryset(report_type, filters)
        if not qs.exists():
            raise Http404('No property records match the current filters.')

        equipment_list = list(
            qs.select_related(
                'assigned_user', 'assigned_user__role', 'assigned_user__department',
                'department', 'category', 'property_status', 'fund_source',
            )
        )

        from core.signatories import resolve_par_download_supply_officer
        supply_officer = resolve_par_download_supply_officer(
            request.user,
            request.GET.get('supply_officer'),
        )

        from .inventory_excel import generate_inventory_excel

        output, filename = generate_inventory_excel(
            equipment_list,
            report_type=report_type,
            supply_officer=supply_officer,
        )
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    qs, columns = get_report_data(report_type, ctx['filters'])

    wb = Workbook()
    ws = wb.active
    report_label = dict(REPORT_TYPES).get(report_type, 'Report')
    ws.title = _excel_sheet_title(report_label)

    ws.append(['PROPERTY MANAGEMENT SYSTEM'])
    ws.append([report_label])
    ws.append([f'Generated: {date.today()}'])
    ws.append([f'Total Records: {ctx["count"]}'])
    ws.append([])

    header_fill = PatternFill(fill_type='solid', fgColor='1E3A5F')
    section_fill = PatternFill(fill_type='solid', fgColor='D6E4F0')
    header_font = Font(bold=True, color='FFFFFF')
    section_font = Font(bold=True, color='1E3A5F')

    def write_header_row():
        ws.append(columns)
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

    if report_type in PROPERTY_REPORT_TYPES:
        for idx, section in enumerate(ctx['sections']):
            if idx > 0:
                ws.append([])
            ws.append([
                f"Office: {section['office']}  |  "
                f"Personnel: {section['personnel']}  |  "
                f"Category: {section['category']}  |  "
                f"Fund Source: {section['fund_source']}  "
                f"({section['count']} item{'s' if section['count'] != 1 else ''})",
            ])
            for cell in ws[ws.max_row]:
                cell.font = section_font
                cell.fill = section_fill
            write_header_row()
            for row in section['rows']:
                ws.append(row)
    else:
        write_header_row()
        for obj in qs:
            ws.append(row_for(obj, report_type))

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 45)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{report_type}_{date.today()}.xlsx"'
    return response


@insights_login_required
def export_pdf(request):
    """Export report as PDF using ReportLab."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    ctx = _report_context(request)
    report_type = ctx['selected_type']
    qs, columns = get_report_data(report_type, ctx['filters'])

    report_label = dict(REPORT_TYPES).get(report_type, 'Report')
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm)
    styles = getSampleStyleSheet()
    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading3'],
        fontSize=9,
        textColor=colors.HexColor('#080'),
        spaceAfter=4,
        spaceBefore=8,
    )
    elements = []

    elements.append(Paragraph('PROPERTY MANAGEMENT SYSTEM', styles['Title']))
    elements.append(Paragraph(report_label, styles['Heading2']))
    elements.append(Paragraph(
        f'Generated: {date.today()} &nbsp;|&nbsp; Total Records: {ctx["count"]}',
        styles['Normal'],
    ))
    elements.append(Spacer(1, 0.4*cm))

    table_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#080')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F4F8')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 3),
    ])

    if report_type in PROPERTY_REPORT_TYPES:
        if not ctx['sections']:
            elements.append(Paragraph('No data available.', styles['Normal']))
        else:
            for section in ctx['sections']:
                header_text = (
                    f"<b>Office:</b> {section['office']} &nbsp;|&nbsp; "
                    f"<b>Personnel:</b> {section['personnel']} &nbsp;|&nbsp; "
                    f"<b>Category:</b> {section['category']} &nbsp;|&nbsp; "
                    f"<b>Fund Source:</b> {section['fund_source']} "
                    f"({section['count']} item{'s' if section['count'] != 1 else ''})"
                )
                elements.append(Paragraph(header_text, section_style))
                data = [columns] + [[str(v) for v in row] for row in section['rows']]
                table = Table(data, repeatRows=1)
                table.setStyle(table_style)
                elements.append(table)
                elements.append(Spacer(1, 0.2*cm))
    else:
        data = [columns]
        for obj in qs:
            data.append([str(v) for v in row_for(obj, report_type)])
        if len(data) > 1:
            table = Table(data, repeatRows=1)
            table.setStyle(table_style)
            elements.append(table)
        else:
            elements.append(Paragraph('No data available.', styles['Normal']))

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{report_type}_{date.today()}.pdf"'
    return response

@insights_login_required
def export_inventory_pdf(request):
    """Export COA inventory report as PDF."""
    ctx = _report_context(request)
    report_type = ctx['selected_type']
    if report_type not in PROPERTY_REPORT_TYPES:
        raise Http404('COA inventory PDF is only available for property reports.')

    filters = ctx['filters']
    qs = _property_queryset(report_type, filters)
    if not qs.exists():
        raise Http404('No property records match the current filters.')

    equipment_list = list(
        qs.select_related(
            'assigned_user', 'assigned_user__role', 'assigned_user__department',
            'department', 'category', 'property_status', 'fund_source',
        )
    )

    from core.signatories import resolve_par_download_supply_officer
    from .inventory_pdf import generate_inventory_pdf

    supply_officer = resolve_par_download_supply_officer(
        request.user,
        request.GET.get('supply_officer'),
    )
    buffer, filename = generate_inventory_pdf(
        equipment_list,
        report_type=report_type,
        supply_officer=supply_officer,
    )
    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _par_report_export_setup(request):
    """Shared queryset and signatory context for PAR report exports."""
    from core.signatories import end_user_position, parse_par_export_signature, resolve_par_download_supply_officer, resolve_par_received_by_user
    from equipment.coa import CLASSIFICATION_PPE
    from equipment.par import ensure_par_for_queryset, par_entity_name
    from .models import ParExportSignatory

    selected_type = request.GET.get('type', 'property_all')
    if selected_type not in PROPERTY_REPORT_TYPES:
        raise Http404('PAR format is only available for property reports.')

    filters = _parse_filters(request.GET)
    qs = _property_queryset(selected_type, filters)
    if not qs.exists():
        raise Http404('No property records match the current filters.')

    equipment_list = list(qs.select_related('assigned_user', 'assigned_user__role', 'assigned_user__department'))
    end_user = resolve_par_received_by_user(
        equipment_list=equipment_list,
        personnel_id=filters.get('personnel_id'),
    ) or request.user
    entity_name = par_entity_name(equipment_list=equipment_list, end_user=end_user)
    received_by_display_name = end_user.get_full_name() or end_user.username

    received_by_name = (request.GET.get('received_by_name') or '').strip()
    if received_by_name:
        try:
            signatory_overrides, supply_officer = parse_par_export_signature(
                request, end_user=end_user,
            )
        except ValueError as exc:
            raise Http404(str(exc)) from exc
    else:
        signatory_overrides = {
            'received_by_name': received_by_display_name,
            'received_by_position': end_user_position(end_user),
        }
        supply_officer = resolve_par_download_supply_officer(
            request.user,
            request.GET.get('supply_officer'),
        )

    ensure_par_for_queryset(
        qs.filter(classification=CLASSIFICATION_PPE),
        created_by=request.user,
    )

    if received_by_name:
        ParExportSignatory.objects.create(
            signature_over_printed_name=signatory_overrides['received_by_name'],
            position_office=signatory_overrides['received_by_position'],
            supply_officer=supply_officer,
            supply_officer_position=signatory_overrides['received_from_position'],
            exported_by=request.user,
            report_type=selected_type,
            item_count=len(equipment_list),
            filter_department_id=filters.get('dept_id') or '',
            filter_personnel_id=filters.get('personnel_id') or '',
        )

    return {
        'equipment_list': equipment_list,
        'entity_name': entity_name,
        'end_user': end_user,
        'signatory_overrides': signatory_overrides,
        'supply_officer': supply_officer,
    }


@insights_login_required
def export_par(request):
    """Export the current property report in the official PAR format (Annex A.1).

    Signatories are resolved automatically from the database. Optional GET
    params may still override fields; saves a ParExportSignatory when provided.
    """
    from .coa_pdf import generate_par_pdf_from_equipment

    ctx = _par_report_export_setup(request)
    buffer, filename = generate_par_pdf_from_equipment(
        ctx['equipment_list'],
        supply_officer=ctx['supply_officer'],
        end_user=ctx['end_user'],
        signatory_overrides=ctx['signatory_overrides'],
    )
    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@insights_login_required
def export_par_excel(request):
    """Export the current property report in PAR Excel format (Appendix 71)."""
    from .par_excel import generate_par_excel

    ctx = _par_report_export_setup(request)
    output, filename = generate_par_excel(
        ctx['equipment_list'],
        supply_officer=ctx['supply_officer'],
        end_user=ctx['end_user'],
        signatory_overrides=ctx['signatory_overrides'],
    )
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
