"""ICS (Appendix 59) Excel layout — mirrors the official ICS PDF form."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from core.models import SiteSettings
from core.signatories import resolve_supply_officer

from .coa_pdf import (
    _fmt_amount,
    _group_ics_by_entity,
    _ics_compute_body_fill,
    _ics_fund_cluster_for_group,
    _ics_line_items_from_group,
    _ics_max_items_per_page,
    _ics_numbers_for_group,
    _ics_supplier_block,
    _ics_usable_dimensions,
    _manila_today,
    _par_sig_date_display,
    _user_office,
)

_SIDE = Side(style='thin', color='000000')
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)

_COLS = 8
_ROW_HEADER = 22
_ROW_TABLE = 18
_ROW_SIG_HEADER = 20
_ROW_SIG_VALUE = 22
_ROW_SIG_NOTE = 16
_ROW_SIG_DATE_VAL = 16


def _set_height(ws, row, height):
    ws.row_dimensions[row].height = height


def _border_row(ws, row, cols=_COLS):
    for col in range(1, cols + 1):
        ws.cell(row, col).border = _BORDER


def _merge_row(ws, row, col_start, col_end, value, *, font, align, height):
    if col_start < col_end:
        ws.merge_cells(
            start_row=row, start_column=col_start,
            end_row=row, end_column=col_end,
        )
    cell = ws.cell(row, col_start, value)
    cell.font = font
    cell.alignment = align
    _border_row(ws, row)
    _set_height(ws, row, height)


def _ics_signatories(primary, site):
    """Resolve ICS signature block names and positions."""
    received_by = primary.received_by
    received_from = primary.received_from or resolve_supply_officer(None)
    return {
        'by_name': (
            (received_by.get_full_name() if received_by else '')
            or site.ics_received_by_name
        ),
        'from_name': (
            (received_from.get_full_name() if received_from else '')
            or site.ics_received_from_name
        ),
        'by_position': (
            primary.received_by_position
            or _user_office(received_by)
            or site.ics_received_by_position
        ),
        'from_position': (
            primary.received_from_position
            or _user_office(received_from)
            or site.ics_received_from_position
        ),
    }


def _write_sig_row(
    ws,
    row,
    left,
    right,
    *,
    left_font,
    right_font=None,
    left_align='left',
    right_align='left',
    height=_ROW_TABLE,
    pad_left=False,
    left_cols=4,
    right_cols=4,
):
    """Two-column signature row split across eight columns."""
    if right_font is None:
        right_font = left_font
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=left_cols)
    ws.merge_cells(
        start_row=row, start_column=left_cols + 1, end_row=row, end_column=_COLS,
    )
    lc = ws.cell(row, 1, left)
    rc = ws.cell(row, left_cols + 1, right)
    lc.font = left_font
    rc.font = right_font
    lc.alignment = Alignment(
        horizontal=left_align, vertical='center', wrap_text=True,
        indent=1 if pad_left else 0,
    )
    rc.alignment = Alignment(
        horizontal=right_align, vertical='center', wrap_text=True,
        indent=1 if pad_left else 0,
    )
    _border_row(ws, row)
    _set_height(ws, row, height)


def _write_item_row(ws, row, values, *, header=False, subheader=False):
    """Write one ICS item row — description spans columns C-D."""
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.border = _BORDER
        cell.font = Font(bold=header or subheader, size=9)
        if header or subheader:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        elif col in (1, 2, 7, 8):
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        elif col in (5, 6):
            cell.alignment = Alignment(horizontal='right', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            if not header and len(str(value or '')) > 22:
                cell.font = Font(size=8)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
    _set_height(ws, row, _ROW_TABLE)


def _write_table_header(ws, row):
    """Two-row ICS item table header."""
    header_row = row
    sub_row = row + 1
    headers = [
        'Qty.', 'Unit', 'Description', '', 'Amount', '', 'Inventory Item No.', 'Estimated Useful Life',
    ]
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, value)
        cell.border = _BORDER
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.cell(sub_row, 5, 'Unit Cost').font = Font(bold=True, size=9)
    ws.cell(sub_row, 5).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(sub_row, 5).border = _BORDER
    ws.cell(sub_row, 6, 'Total Cost').font = Font(bold=True, size=9)
    ws.cell(sub_row, 6).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(sub_row, 6).border = _BORDER
    ws.cell(sub_row, 3, '').border = _BORDER
    ws.cell(sub_row, 4, '').border = _BORDER

    ws.merge_cells(start_row=header_row, start_column=1, end_row=sub_row, end_column=1)
    ws.merge_cells(start_row=header_row, start_column=2, end_row=sub_row, end_column=2)
    ws.merge_cells(start_row=header_row, start_column=3, end_row=header_row, end_column=4)
    ws.merge_cells(start_row=sub_row, start_column=3, end_row=sub_row, end_column=4)
    ws.merge_cells(start_row=header_row, start_column=5, end_row=header_row, end_column=6)
    ws.merge_cells(start_row=header_row, start_column=7, end_row=sub_row, end_column=7)
    ws.merge_cells(start_row=header_row, start_column=8, end_row=sub_row, end_column=8)

    _set_height(ws, header_row, _ROW_TABLE)
    _set_height(ws, sub_row, _ROW_TABLE)


def _ics_blank_row_count(item_count):
    _, usable_h = _ics_usable_dimensions()
    blank_rows, _ = _ics_compute_body_fill(item_count, usable_h)
    return blank_rows


def _setup_ics_columns(ws, line_items=None):
    ws.column_dimensions['A'].width = 7
    ws.column_dimensions['B'].width = 7
    desc_width = 14
    inv_width = 14
    if line_items:
        max_inv = max(
            (len(str(getattr(item, 'inventory_item_no', '') or '')) for item in line_items),
            default=0,
        )
        if max_inv > 18:
            inv_width = min(24, max(14, max_inv * 0.75 + 2))
            desc_width = max(12, 14 - (inv_width - 14) * 0.5)
    ws.column_dimensions['C'].width = desc_width
    ws.column_dimensions['D'].width = desc_width
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = inv_width
    ws.column_dimensions['H'].width = 10


def _excel_sheet_title(ics_no, index=0):
    """Sanitize ICS number for use as an Excel worksheet title (max 31 chars)."""
    invalid = set(r'\/*?:[]')
    cleaned = ''.join('-' if ch in invalid else ch for ch in (ics_no or 'ICS')).strip()
    title = (cleaned or 'ICS')[:31]
    if index:
        suffix = f'-{index}'
        title = f'{title[:31 - len(suffix)]}{suffix}'
    return title or 'ICS'


def _write_supplier_field(ws, row, label, value):
    """Right-aligned supplier block field inside the header box."""
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=_COLS)
    cell = ws.cell(row, 5, f'{label} {value or ""}'.strip())
    cell.font = Font(size=9, bold=True)
    cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    _border_row(ws, row)
    _set_height(ws, row, 16)


def _write_ics_form_sheet(
    ws,
    *,
    site,
    ics_no,
    entity_name,
    fund_cluster,
    supplier_block,
    line_items,
    signatories,
    sig_date_by,
    sig_date_from=None,
    page_label='Page 1/1',
):
    """Write one Appendix 59 ICS form onto a worksheet."""
    if sig_date_from is None:
        sig_date_from = sig_date_by

    _setup_ics_columns(ws, line_items)

    row = 1
    annex = ws.cell(row, 7, 'Appendix 59')
    annex.font = Font(italic=True, size=9)
    annex.alignment = Alignment(horizontal='right', vertical='bottom')
    circular = ws.cell(row, 8, 'COA Circular 2022-004')
    circular.font = Font(italic=True, size=9)
    circular.alignment = Alignment(horizontal='right', vertical='bottom')
    _set_height(ws, row, 15)
    row += 1

    _merge_row(
        ws, row, 1, _COLS, 'INVENTORY CUSTODIAN SLIP',
        font=Font(bold=True, size=12),
        align=Alignment(horizontal='center', vertical='center'),
        height=_ROW_HEADER,
    )
    row += 1

    _write_supplier_field(ws, row, 'Supplier:', supplier_block['supplier'])
    row += 1
    _write_supplier_field(ws, row, 'Address:', supplier_block['address'])
    row += 1
    _write_supplier_field(ws, row, 'P.O. #:', supplier_block['po_number'])
    row += 1
    _write_supplier_field(ws, row, 'Date:', supplier_block['date'])
    row += 1
    _write_supplier_field(ws, row, 'MOP:', supplier_block['mop'])
    row += 1

    _merge_row(
        ws, row, 1, _COLS, f'Entity Name: {entity_name}',
        font=Font(size=10),
        align=Alignment(horizontal='left', vertical='center', indent=1),
        height=18,
    )
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=_COLS)
    fund_cell = ws.cell(row, 1, f'Fund Cluster: {fund_cluster}')
    fund_cell.font = Font(size=10)
    fund_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ics_cell = ws.cell(row, 5, f'ICS No.: {ics_no}')
    ics_cell.font = Font(size=10)
    ics_cell.alignment = Alignment(horizontal='right', vertical='center')
    _border_row(ws, row)
    _set_height(ws, row, 18)
    row += 1

    row += 1
    table_start = row
    _write_table_header(ws, row)
    row += 2

    for item in line_items:
        _write_item_row(ws, row, [
            item.quantity,
            item.unit,
            item.description or '',
            '',
            _fmt_amount(item.estimated_unit_cost),
            _fmt_amount(item.total_cost),
            item.inventory_item_no,
            item.useful_life,
        ])
        row += 1

    for _ in range(_ics_blank_row_count(len(line_items))):
        _write_item_row(ws, row, ['', '', '', '', '', '', '', ''])
        row += 1

    sig_header = Font(bold=True, size=9)
    sig_note = Font(size=7)
    sig_name = Font(bold=True, size=9, underline='single')
    sig_pos = Font(bold=True, size=9, underline='single')
    sig_date_font = Font(bold=True, size=8, underline='single')

    _write_sig_row(
        ws, row, 'Received by:', 'Received from:',
        left_font=sig_header, left_align='left', right_align='left',
        height=_ROW_SIG_HEADER, pad_left=True,
    )
    row += 1

    _write_sig_row(
        ws, row,
        (signatories['by_name'] or '').upper(),
        (signatories['from_name'] or '').upper(),
        left_font=sig_name, left_align='center', right_align='center',
        height=_ROW_SIG_VALUE,
    )
    row += 1

    _write_sig_row(
        ws, row,
        'Signature Over Printed Name',
        'Signature Over Printed Name',
        left_font=sig_note, left_align='center', right_align='center',
        height=_ROW_SIG_NOTE,
    )
    row += 1

    _write_sig_row(
        ws, row,
        (signatories['by_position'] or '').upper(),
        (signatories['from_position'] or '').upper(),
        left_font=sig_pos, left_align='center', right_align='center',
        height=_ROW_SIG_VALUE,
    )
    row += 1

    _write_sig_row(
        ws, row,
        'Position/Office',
        'Position/Office',
        left_font=sig_note, left_align='center', right_align='center',
        height=_ROW_SIG_NOTE,
    )
    row += 1

    _write_sig_row(
        ws, row,
        _par_sig_date_display(sig_date_by),
        _par_sig_date_display(sig_date_from),
        left_font=sig_date_font, left_align='center', right_align='center',
        height=_ROW_SIG_DATE_VAL,
    )
    row += 1

    _write_sig_row(
        ws, row, 'Date', 'Date',
        left_font=sig_note, left_align='center', right_align='center',
        height=_ROW_SIG_NOTE,
    )
    row += 1

    form_end = row - 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=_COLS)
    footer_left = ws.cell(row, 1, 'Inventory Custodian Slip')
    footer_left.font = Font(italic=True, size=8, color='808080')
    footer_left.alignment = Alignment(horizontal='left', vertical='center')
    footer_right = ws.cell(row, 5, page_label)
    footer_right.font = Font(italic=True, size=8, color='808080')
    footer_right.alignment = Alignment(horizontal='right', vertical='center')
    _set_height(ws, row, 14)

    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.4, bottom=0.4)
    ws.print_area = f'A{table_start}:H{form_end}'


def generate_all_ics_excel(ics_list):
    """Build a workbook — same entity name combined on one sheet per form."""
    ics_records = list(ics_list)
    if not ics_records:
        raise ValueError('No ICS records to export.')

    site = SiteSettings.get_solo()
    wb = Workbook()
    wb.remove(wb.active)

    used_titles = set()
    sheet_index = 0
    per_page = _ics_max_items_per_page()
    for group in _group_ics_by_entity(ics_records):
        primary = group[0]
        all_line_items = _ics_line_items_from_group(group)
        ics_no = _ics_numbers_for_group(group) if len(group) > 1 else primary.ics_no
        entity_name = primary.entity_name
        fund_cluster = _ics_fund_cluster_for_group(group)
        supplier_block = _ics_supplier_block(site, primary)
        signatories = _ics_signatories(primary, site)
        sig_date_by = str(primary.date_received_by or primary.date or _manila_today())
        sig_date_from = str(primary.date_received_from or primary.date or _manila_today())
        chunks = [
            all_line_items[i:i + per_page]
            for i in range(0, len(all_line_items), per_page)
        ] or [[]]
        page_total = len(chunks)
        for idx, chunk in enumerate(chunks, start=1):
            sheet_index += 1
            title = _excel_sheet_title(ics_no)
            if title in used_titles:
                title = _excel_sheet_title(ics_no, sheet_index)
            used_titles.add(title)

            ws = wb.create_sheet(title=title)
            _write_ics_form_sheet(
                ws,
                site=site,
                ics_no=ics_no,
                entity_name=entity_name,
                fund_cluster=fund_cluster,
                supplier_block=supplier_block,
                line_items=chunk,
                signatories=signatories,
                sig_date_by=sig_date_by,
                sig_date_from=sig_date_from,
                page_label=f'Page {idx}/{page_total}',
            )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    item_count = sum(ics.items.count() for ics in ics_records)
    return output, f'ICS_All_{item_count}_records_{date.today()}.xlsx'
